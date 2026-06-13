import json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
import sys


def set_heading_color(para, hex_color):
    for run in para.runs:
        run.font.color.rgb = RGBColor.from_string(hex_color)

def add_divider(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CADCFC")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)

here = Path(__file__).resolve().parent
cfg = here.parent / "cache" / "autopsy.json"
if not cfg.exists():
    print(f"ERROR: expected file not found: {cfg}")
    print("Make sure 'cache/autopsy.json' exists relative to the project root.")
    sys.exit(1)
with cfg.open("r", encoding="utf-8") as f:
    data = json.load(f)

doc = Document()

# ── Page margins ────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.1)
    section.right_margin  = Inches(1.1)

# ── Title ────────────────────────────────────────────────────
brand = data["per_ad"][0]["ad_id"].split("_")[0].upper()
title = doc.add_heading(f"Ad Autopsy: {brand}", 0)
title.runs[0].font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
title.runs[0].font.size = Pt(28)

sub = doc.add_paragraph(
    f"{len(data['per_ad'])} ads analysed  ·  "
    f"{len(data['patterns'])} patterns found  ·  1 gap identified"
)
sub.runs[0].font.color.rgb = RGBColor(0x8A, 0x9B, 0xB8)
sub.runs[0].font.size = Pt(11)
sub.runs[0].italic = True
doc.add_paragraph()

# ── Per-ad section ───────────────────────────────────────────
h = doc.add_heading("Per-Ad Analysis", level=1)
set_heading_color(h, "1E2761")

for i, ad in enumerate(data["per_ad"]):
    add_divider(doc)
    h2 = doc.add_heading(f"Ad {i+1}: {ad['ad_id']}", level=2)
    set_heading_color(h2, "2A3A8C")

    # Hook
    p = doc.add_paragraph()
    r = p.add_run("Hook (t={}s):  ".format(ad["hook"]["t"]))
    r.bold = True
    r.font.color.rgb = RGBColor(0xF9, 0x61, 0x67)
    p.add_run(ad["hook"]["desc"])

    # Emotional arc
    p = doc.add_paragraph()
    r = p.add_run("Emotional Arc:  ")
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
    p.add_run(ad["emotional_arc"])

    # Visual pattern
    p = doc.add_paragraph()
    r = p.add_run("Visual Pattern:  ")
    r.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
    p.add_run(ad["visual_pattern"])

    # Claims table
    doc.add_paragraph()
    h3 = doc.add_heading("Claims", level=3)
    set_heading_color(h3, "8A9BB8")

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Timestamp"
    hdr[1].text = "Claim"
    for cell in hdr:
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1E2761")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    for claim in ad["claims"]:
        row = table.add_row().cells
        row[0].text = f"t={claim['t']}s"
        row[1].text = claim["text"]

    doc.add_paragraph()

# ── Patterns ─────────────────────────────────────────────────
add_divider(doc)
h = doc.add_heading("Cross-Ad Patterns", level=1)
set_heading_color(h, "1E2761")

for i, pattern in enumerate(data["patterns"]):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(pattern["insight"])
    r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)

doc.add_paragraph()

# ── The Gap ───────────────────────────────────────────────────
add_divider(doc)
h = doc.add_heading("The Gap", level=1)
set_heading_color(h, "F96167")

gap_p = doc.add_paragraph()
gap_p.paragraph_format.left_indent = Inches(0.3)
gap_run = gap_p.add_run(f'"{data["gap"]}"')
gap_run.italic = True
gap_run.font.size = Pt(13)
gap_run.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)

doc.add_paragraph()

# ── Counter-creative ─────────────────────────────────────────
add_divider(doc)
h = doc.add_heading("Counter-Creative", level=1)
set_heading_color(h, "1E2761")

p = doc.add_paragraph()
r = p.add_run("Angle:  ")
r.bold = True
r.font.color.rgb = RGBColor(0xF9, 0x61, 0x67)
p.add_run(data["counter_creative"]["angle"])

doc.add_paragraph()
h3 = doc.add_heading("Script", level=3)
set_heading_color(h3, "8A9BB8")

script_p = doc.add_paragraph()
script_p.paragraph_format.left_indent = Inches(0.3)
script_run = script_p.add_run(data["counter_creative"]["script"])
script_run.font.size = Pt(11)
script_run.italic = True

doc.save("ad_autopsy.docx")
print("ad_autopsy.docx saved ✓")